import os
import re
import tempfile
import mimetypes
import uuid
from pathlib import Path
from urllib.parse import urlparse

import httpx
from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    UploadFile,
    File,
    Form,
)
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from starlette.background import BackgroundTask

from ..database import get_db
from ..deps import get_current_user

from ..models import (
    ResumeVersion,
    ResumeAnalysis,
    MatchAnalysis,
    JobApplication,
)

from ..schemas import (
    ResumeCreate,
    ResumeOut,
)

from ..services.resume_parser import (
    extract_text_from_pdf,
)


# ============================================================
# ROUTER
# ============================================================

router = APIRouter(
    prefix="/resumes",
    tags=["Resumes"],
)


# ============================================================
# CONFIGURATION
# ============================================================

UPLOAD_DIR = Path("uploads")
BLOB_API_URL = "https://blob.vercel-storage.com"
BLOB_API_VERSION = "11"
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "4"))
if MAX_UPLOAD_MB < 1:
    raise ValueError("MAX_UPLOAD_MB must be at least 1.")
MAX_UPLOAD_BYTES = min(
    MAX_UPLOAD_MB,
    4 if os.getenv("VERCEL") else MAX_UPLOAD_MB,
) * 1024 * 1024

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
}


# ============================================================
# RESUME SKILL ALIASES
# ============================================================
#
# The left side is the canonical skill saved in the database.
#
# Example:
#
#   react.js
#   reactjs
#   React
#
# are all stored as:
#
#   react
#
# ============================================================

RESUME_SKILL_ALIASES: dict[str, list[str]] = {

    # --------------------------------------------------------
    # Programming Languages
    # --------------------------------------------------------

    "python": [
        "python",
    ],

    "java": [
        "java",
    ],

    "javascript": [
        "javascript",
        "java script",
        "js",
    ],

    "typescript": [
        "typescript",
        "type script",
        "ts",
    ],

    "c++": [
        "c++",
    ],

    "c#": [
        "c#",
        "c sharp",
    ],

    "go": [
        "golang",
        "go language",
    ],

    "php": [
        "php",
    ],

    "ruby": [
        "ruby",
    ],

    # --------------------------------------------------------
    # Frontend
    # --------------------------------------------------------

    "html": [
        "html",
        "html5",
    ],

    "css": [
        "css",
        "css3",
    ],

    "react": [
        "react",
        "react.js",
        "reactjs",
    ],

    "next.js": [
        "next.js",
        "nextjs",
        "next js",
    ],

    "angular": [
        "angular",
        "angularjs",
    ],

    "vue": [
        "vue",
        "vue.js",
        "vuejs",
    ],

    "tailwind": [
        "tailwind",
        "tailwind css",
    ],

    "bootstrap": [
        "bootstrap",
    ],

    # --------------------------------------------------------
    # Backend
    # --------------------------------------------------------

    "node.js": [
        "node.js",
        "nodejs",
        "node js",
    ],

    "express": [
        "express",
        "express.js",
        "expressjs",
    ],

    "fastapi": [
        "fastapi",
        "fast api",
    ],

    "django": [
        "django",
    ],

    "flask": [
        "flask",
    ],

    "spring": [
        "spring framework",
    ],

    "spring boot": [
        "spring boot",
        "springboot",
    ],

    # --------------------------------------------------------
    # Databases
    # --------------------------------------------------------

    "sql": [
        "sql",
        "structured query language",
    ],

    "mysql": [
        "mysql",
        "my sql",
    ],

    "postgresql": [
        "postgresql",
        "postgres",
        "postgre sql",
    ],

    "mongodb": [
        "mongodb",
        "mongo db",
    ],

    "redis": [
        "redis",
    ],

    "oracle": [
        "oracle database",
        "oracle db",
    ],

    # --------------------------------------------------------
    # Data Science
    # --------------------------------------------------------

    "pandas": [
        "pandas",
    ],

    "numpy": [
        "numpy",
    ],

    "scikit-learn": [
        "scikit-learn",
        "scikit learn",
        "sklearn",
    ],

    "tensorflow": [
        "tensorflow",
    ],

    "pytorch": [
        "pytorch",
    ],

    "machine learning": [
        "machine learning",
        "machine-learning",
    ],

    "deep learning": [
        "deep learning",
        "deep-learning",
    ],

    "data science": [
        "data science",
        "data scientist",
    ],

    "data analysis": [
        "data analysis",
        "data analytics",
        "data analyst",
    ],

    "data visualization": [
        "data visualization",
        "data visualisation",
    ],

    "power bi": [
        "power bi",
        "powerbi",
    ],

    "tableau": [
        "tableau",
    ],

    "excel": [
        "excel",
        "microsoft excel",
    ],

    # --------------------------------------------------------
    # APIs
    # --------------------------------------------------------

    "rest api": [
        "rest api",
        "rest apis",
        "restful api",
        "restful apis",
    ],

    "graphql": [
        "graphql",
    ],

    "api": [
        "api",
        "apis",
        "application programming interface",
    ],

    # --------------------------------------------------------
    # DevOps / Cloud
    # --------------------------------------------------------

    "git": [
        "git",
    ],

    "github": [
        "github",
        "git hub",
    ],

    "gitlab": [
        "gitlab",
        "git lab",
    ],

    "docker": [
        "docker",
    ],

    "kubernetes": [
        "kubernetes",
        "k8s",
    ],

    "aws": [
        "aws",
        "amazon web services",
    ],

    "azure": [
        "azure",
        "microsoft azure",
    ],

    "gcp": [
        "gcp",
        "google cloud",
        "google cloud platform",
    ],

    "linux": [
        "linux",
    ],

    "jenkins": [
        "jenkins",
    ],

    "terraform": [
        "terraform",
    ],

    "ci/cd": [
        "ci/cd",
        "cicd",
        "continuous integration",
        "continuous delivery",
    ],

    # --------------------------------------------------------
    # Development Tools
    # --------------------------------------------------------

    "postman": [
        "postman",
    ],

    "swagger": [
        "swagger",
        "openapi",
        "open api",
    ],

    "pytest": [
        "pytest",
    ],

    # --------------------------------------------------------
    # Concepts
    # --------------------------------------------------------

    "oop": [
        "oop",
        "object oriented programming",
        "object-oriented programming",
    ],

    "agile": [
        "agile",
        "scrum",
    ],

    "rest": [
        "rest",
        "restful",
    ],
}


# ============================================================
# HELPER FUNCTIONS
# ============================================================


def normalize_text(
    text: str | None,
) -> str:
    """
    Normalize text before skill matching.

    Example:

        "Python   Developer"
            ->
        "python developer"
    """

    if not text:
        return ""

    text = text.lower()

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


def contains_term(
    text: str,
    term: str,
) -> bool:
    """
    Safely check whether a term exists.

    This prevents:

        java

    from matching:

        javascript
    """

    normalized_text = normalize_text(
        text
    )

    normalized_term = normalize_text(
        term
    )

    if (
        not normalized_text
        or not normalized_term
    ):
        return False

    pattern = (
        r"(?<![a-z0-9])"
        + re.escape(normalized_term)
        + r"(?![a-z0-9])"
    )

    return (
        re.search(
            pattern,
            normalized_text,
        )
        is not None
    )


def extract_resume_skills(
    text: str,
) -> list[str]:
    """
    Extract technical skills from resume text.

    Returns canonical skill names.
    """

    normalized_text = normalize_text(
        text
    )

    detected = set()

    if not normalized_text:
        return []

    for (
        canonical_skill,
        aliases,
    ) in RESUME_SKILL_ALIASES.items():

        for alias in aliases:

            if contains_term(
                normalized_text,
                alias,
            ):
                detected.add(
                    canonical_skill
                )

                break

    return sorted(
        detected
    )


def clean_manual_skills(
    skills: str,
) -> list[str]:
    """
    Clean manually entered skills.
    """

    if not skills:
        return []

    result = []

    for skill in skills.split(","):

        cleaned = normalize_text(
            skill
        )

        if cleaned:
            result.append(
                cleaned
            )

    return sorted(
        set(result)
    )


def extract_docx_text(
    file_path: Path,
) -> str:
    """
    Extract text from DOCX paragraphs
    and tables.
    """

    from docx import Document

    document = Document(
        str(file_path)
    )

    parts = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_values.append(
                        cell_text
                    )

            if row_values:
                parts.append(
                    " ".join(row_values)
                )

    return "\n".join(
        parts
    ).strip()


def extract_resume_text(
    file_path: Path,
) -> str:
    """
    Extract readable text depending
    on the uploaded file type.
    """

    extension = (
        file_path.suffix.lower()
    )

    if extension == ".pdf":

        return extract_text_from_pdf(
            str(file_path)
        )

    if extension == ".docx":

        return extract_docx_text(
            file_path
        )

    if extension == ".doc":

        raise HTTPException(
            status_code=400,
            detail=(
                "DOC files are not currently "
                "supported for automatic AI "
                "analysis. Please convert the "
                "resume to PDF or DOCX."
            ),
        )

    raise HTTPException(
        status_code=400,
        detail=(
            "Unsupported resume file format."
        ),
    )


def build_resume_summary(
    skills: list[str],
    extracted_text: str,
) -> str:
    """
    Build a deterministic resume summary.
    """

    preview = " ".join(
        extracted_text.split()
    )

    if len(preview) > 300:
        preview = (
            preview[:300].rstrip()
            + "..."
        )

    summary = (
        f"Detected {len(skills)} "
        f"technical skills from the resume."
    )

    if skills:

        summary += (
            " Skills: "
            + ", ".join(skills)
            + "."
        )

    if preview:

        summary += (
            f" Resume preview: {preview}"
        )

    return summary


# ============================================================
# GET OWNED RESUME
# ============================================================


def get_owned_resume(
    db: Session,
    resume_id: int,
    user,
) -> ResumeVersion:

    resume = (
        db.query(ResumeVersion)
        .filter(
            ResumeVersion.id == resume_id,
            ResumeVersion.user_id == user.id,
        )
        .first()
    )

    if not resume:

        raise HTTPException(
            status_code=404,
            detail="Resume not found.",
        )

    return resume


# ============================================================
# VALIDATE FILE PATH
# ============================================================


def validate_file_path(
    file_path: str,
) -> Path:

    upload_root = (
        UPLOAD_DIR.resolve()
    )

    path = Path(
        file_path
    ).resolve()

    try:

        path.relative_to(
            upload_root
        )

    except ValueError:

        raise HTTPException(
            status_code=400,
            detail="Invalid resume file path.",
        )

    return path


def _blob_token() -> str:
    token = os.getenv("BLOB_READ_WRITE_TOKEN", "")
    if not token:
        raise HTTPException(
            status_code=503,
            detail=(
                "Resume storage is not configured. "
                "Connect a Vercel Blob store and set "
                "BLOB_READ_WRITE_TOKEN."
            ),
        )
    return token


def _stored_blob_url(file_path: str, user_id: int) -> str:
    prefix = "vercel-blob:"
    if not file_path.startswith(prefix):
        raise HTTPException(
            status_code=400,
            detail="Invalid stored resume location.",
        )

    url = file_path[len(prefix):]
    parsed = urlparse(url)
    if (
        parsed.scheme != "https"
        or not parsed.hostname
        or not parsed.hostname.endswith(".private.blob.vercel-storage.com")
        or parsed.username
        or parsed.password
        or parsed.query
        or parsed.fragment
        or not parsed.path.startswith(f"/resumes/{user_id}/")
    ):
        raise HTTPException(
            status_code=400,
            detail="Invalid stored resume location.",
        )
    return url


async def upload_blob(
    pathname: str,
    content: bytes,
    content_type: str,
    user_id: int,
) -> str:
    token = _blob_token()
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            response = await client.put(
                BLOB_API_URL,
                params={"pathname": pathname},
                content=content,
                headers={
                    "authorization": f"Bearer {token}",
                    "x-api-version": BLOB_API_VERSION,
                    "x-content-type": content_type,
                    "x-vercel-blob-access": "private",
                    "x-add-random-suffix": "0",
                },
            )
    except httpx.HTTPError as exc:
        raise HTTPException(
            status_code=502,
            detail="Could not connect to resume storage.",
        ) from exc

    if response.status_code >= 300:
        raise HTTPException(
            status_code=502,
            detail=(
                "Resume storage rejected the upload "
                f"(HTTP {response.status_code})."
            ),
        )

    try:
        blob_url = response.json().get("url")
    except (ValueError, KeyError, TypeError) as exc:
        raise HTTPException(
            status_code=502,
            detail="Resume storage returned an invalid upload response.",
        ) from exc
    if not isinstance(blob_url, str):
        raise HTTPException(
            status_code=502,
            detail="Resume storage returned an invalid upload response.",
        )

    return _stored_blob_url(
        f"vercel-blob:{blob_url}",
        user_id,
    )


def download_blob_to_temp(blob_url: str, user_id: int, filename: str) -> Path:
    validated_url = _stored_blob_url(
        f"vercel-blob:{blob_url}",
        user_id,
    )
    token = _blob_token()
    try:
        response = httpx.get(
            validated_url,
            headers={"authorization": f"Bearer {token}"},
            timeout=30,
            follow_redirects=True,
        )
    except httpx.HTTPError as exc:
        raise HTTPException(
            status_code=502,
            detail="Could not retrieve the stored resume.",
        ) from exc

    if response.status_code == 404:
        raise HTTPException(
            status_code=404,
            detail="The stored resume file was not found.",
        )
    if response.status_code >= 300:
        raise HTTPException(
            status_code=502,
            detail=(
                "Resume storage could not retrieve the file "
                f"(HTTP {response.status_code})."
            ),
        )

    suffix = Path(filename).suffix or ".bin"
    with tempfile.NamedTemporaryFile(
        prefix="resume_",
        suffix=suffix,
        delete=False,
    ) as temp_file:
        temp_file.write(response.content)
        return Path(temp_file.name)


def delete_blob(blob_url: str, user_id: int) -> None:
    validated_url = _stored_blob_url(
        f"vercel-blob:{blob_url}",
        user_id,
    )
    token = _blob_token()
    try:
        response = httpx.post(
            f"{BLOB_API_URL}/delete",
            json={"urls": [validated_url]},
            headers={
                "authorization": f"Bearer {token}",
                "x-api-version": BLOB_API_VERSION,
                "content-type": "application/json",
            },
            timeout=30,
        )
    except httpx.HTTPError as exc:
        raise HTTPException(
            status_code=502,
            detail="Could not connect to resume storage to delete the file.",
        ) from exc

    if response.status_code >= 300:
        raise HTTPException(
            status_code=502,
            detail=(
                "Resume storage could not delete the file "
                f"(HTTP {response.status_code})."
            ),
        )


# ============================================================
# GET ALL RESUMES
# ============================================================


@router.get(
    "",
    response_model=list[ResumeOut],
)
def list_resumes(
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):

    return (
        db.query(ResumeVersion)
        .filter(
            ResumeVersion.user_id
            == user.id
        )
        .order_by(
            ResumeVersion.created_at.desc()
        )
        .all()
    )


# ============================================================
# CREATE RESUME WITHOUT FILE
# ============================================================


@router.post(
    "",
    response_model=ResumeOut,
)
def create_resume(
    data: ResumeCreate,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):

    resume = ResumeVersion(
        user_id=user.id,
        name=data.name.strip(),
        skills=[
            normalize_text(skill)
            for skill in data.skills
            if normalize_text(skill)
        ],
    )

    db.add(resume)

    db.commit()

    db.refresh(resume)

    return resume


# ============================================================
# UPLOAD RESUME
# ============================================================


@router.post(
    "/upload",
    response_model=ResumeOut,
)
async def upload_resume(
    name: str = Form(""),
    skills: str = Form(""),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):

    original_name = (
        file.filename
        or "resume"
    )

    extension = (
        Path(original_name)
        .suffix
        .lower()
    )

    # --------------------------------------------------------
    # Validate extension
    # --------------------------------------------------------

    if extension not in ALLOWED_EXTENSIONS:

        raise HTTPException(
            status_code=400,
            detail=(
                "Only PDF, DOC and DOCX "
                "files are allowed."
            ),
        )

    # --------------------------------------------------------
    # Read file
    # --------------------------------------------------------

    try:
        content = await file.read(MAX_UPLOAD_BYTES + 1)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Unable to read the uploaded resume file.",
        ) from exc

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded resume file is empty.",
        )
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=(
                "Resume files must be no larger than "
                f"{MAX_UPLOAD_BYTES // (1024 * 1024)} MB."
            ),
        )

    # --------------------------------------------------------
    # Save file
    # --------------------------------------------------------

    stored_filename = (
        f"{uuid.uuid4().hex}"
        f"{extension}"
    )
    blob_url: str | None = None
    if os.getenv("BLOB_READ_WRITE_TOKEN"):
        blob_url = await upload_blob(
            f"resumes/{user.id}/{stored_filename}",
            content,
            mimetypes.guess_type(original_name)[0]
            or "application/octet-stream",
            user.id,
        )
        file_path = f"vercel-blob:{blob_url}"
    elif os.getenv("VERCEL"):
        _blob_token()
    else:
        UPLOAD_DIR.mkdir(
            parents=True,
            exist_ok=True,
        )
        local_path = UPLOAD_DIR / stored_filename
        try:
            local_path.write_bytes(content)
        except OSError as exc:
            raise HTTPException(
                status_code=500,
                detail="Unable to save the uploaded resume.",
            ) from exc
        file_path = str(local_path)

    # --------------------------------------------------------
    # Manual skills
    # --------------------------------------------------------

    manual_skills = (
        clean_manual_skills(
            skills
        )
    )

    # --------------------------------------------------------
    # Create database record
    # --------------------------------------------------------

    resume = ResumeVersion(
        user_id=user.id,
        name=(
            name.strip()
            or Path(
                original_name
            ).stem
        ),
        file_name=original_name,
        file_path=str(file_path),
        skills=manual_skills,
    )

    try:

        db.add(resume)

        db.commit()

        db.refresh(resume)

    except Exception as exc:
        db.rollback()
        if blob_url:
            delete_blob(blob_url, user.id)
        else:
            Path(file_path).unlink(missing_ok=True)
        raise HTTPException(
            status_code=500,
            detail="Unable to save the resume record.",
        ) from exc

    return resume


# ============================================================
# ANALYZE EXISTING RESUME
# ============================================================
#
# POST:
#
# /api/resumes/{resume_id}/analyze
#
# This endpoint:
#
# 1. Gets the uploaded resume
# 2. Extracts PDF/DOCX text
# 3. Detects technical skills
# 4. Merges manual skills
# 5. Saves skills to ResumeVersion
# 6. Creates/updates ResumeAnalysis
#
# ============================================================


@router.post(
    "/{resume_id}/analyze",
)
def analyze_existing_resume(
    resume_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):

    # --------------------------------------------------------
    # 1. Get owned resume
    # --------------------------------------------------------

    resume = get_owned_resume(
        db,
        resume_id,
        user,
    )

    # --------------------------------------------------------
    # 2. Validate uploaded file
    # --------------------------------------------------------

    if not resume.file_path:

        raise HTTPException(
            status_code=400,
            detail=(
                "This resume does not have "
                "an uploaded file."
            ),
        )

    is_blob_file = resume.file_path.startswith("vercel-blob:")
    if is_blob_file:
        blob_url = _stored_blob_url(
            resume.file_path,
            user.id,
        )
        file_path = download_blob_to_temp(
            blob_url,
            user.id,
            resume.file_name or "resume",
        )
    else:
        file_path = validate_file_path(resume.file_path)

    try:
        if not file_path.is_file():
            raise HTTPException(
                status_code=404,
                detail="The resume file is missing from storage.",
            )

        # --------------------------------------------------------
        # 3. Extract text
        # --------------------------------------------------------
        try:
            extracted_text = extract_resume_text(file_path)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail=f"Unable to extract text from resume: {exc}",
            ) from exc
    finally:
        if is_blob_file:
            file_path.unlink(missing_ok=True)

    extracted_text = (
        extracted_text.strip()
    )

    # --------------------------------------------------------
    # 4. Validate extracted text
    # --------------------------------------------------------

    if not extracted_text:

        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text could be "
                "extracted from this resume. "
                "If this is a scanned PDF, "
                "OCR is required."
            ),
        )

    # --------------------------------------------------------
    # 5. Automatically extract skills
    # --------------------------------------------------------

    try:

        detected_skills = (
            extract_resume_skills(
                extracted_text
            )
        )

    except Exception as exc:

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to detect resume "
                f"skills: {exc}"
            ),
        )

    # --------------------------------------------------------
    # 6. Merge manual + automatic skills
    # --------------------------------------------------------
    #
    # If the user manually entered:
    #
    # python, sql
    #
    # and automatic extraction finds:
    #
    # python, pandas, numpy
    #
    # final result:
    #
    # python
    # sql
    # pandas
    # numpy
    #
    # --------------------------------------------------------

    manual_skills = [
        normalize_text(skill)
        for skill in (
            resume.skills
            or []
        )
        if normalize_text(skill)
    ]

    all_skills = sorted(
        set(
            detected_skills
            + manual_skills
        )
    )

    # --------------------------------------------------------
    # 7. Save skills to ResumeVersion
    # --------------------------------------------------------

    resume.skills = all_skills

    # --------------------------------------------------------
    # 8. Build summary
    # --------------------------------------------------------

    summary = build_resume_summary(
        skills=all_skills,
        extracted_text=extracted_text,
    )

    # --------------------------------------------------------
    # 9. Find existing analysis
    # --------------------------------------------------------

    analysis = (
        db.query(
            ResumeAnalysis
        )
        .filter(
            ResumeAnalysis.resume_id
            == resume.id
        )
        .first()
    )

    # --------------------------------------------------------
    # 10. Update existing analysis
    # --------------------------------------------------------

    if analysis:

        analysis.extracted_text = (
            extracted_text
        )

        analysis.detected_skills = (
            all_skills
        )

        analysis.summary = (
            summary
        )

    # --------------------------------------------------------
    # 11. Create new analysis
    # --------------------------------------------------------

    else:

        analysis = ResumeAnalysis(
            resume_id=resume.id,
            extracted_text=(
                extracted_text
            ),
            detected_skills=(
                all_skills
            ),
            summary=summary,
        )

        db.add(analysis)

    # --------------------------------------------------------
    # 12. Save database changes
    # --------------------------------------------------------

    try:

        db.commit()

        db.refresh(resume)

        db.refresh(analysis)

    except Exception as exc:

        db.rollback()

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to save resume "
                f"analysis: {exc}"
            ),
        )

    # --------------------------------------------------------
    # 13. Return result
    # --------------------------------------------------------

    return {
        "id": resume.id,

        "resume_id": resume.id,

        "name": resume.name,

        "file_name": resume.file_name,

        "skills": all_skills,

        "detected_skills": all_skills,

        "automatic_skills": (
            detected_skills
        ),

        "manual_skills": (
            manual_skills
        ),

        "extracted_characters": len(
            extracted_text
        ),

        "summary": summary,

        "analysis_id": analysis.id,

        "created_at": resume.created_at,
    }


# ============================================================
# DELETE RESUME
# ============================================================


@router.delete(
    "/{resume_id}",
)
def delete_resume(
    resume_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):

    resume = get_owned_resume(
        db,
        resume_id,
        user,
    )

    # --------------------------------------------------------
    # Delete ResumeAnalysis
    # --------------------------------------------------------

    db.query(
        ResumeAnalysis
    ).filter(
        ResumeAnalysis.resume_id
        == resume_id
    ).delete(
        synchronize_session=False
    )

    # --------------------------------------------------------
    # Delete MatchAnalysis
    # --------------------------------------------------------

    db.query(
        MatchAnalysis
    ).filter(
        MatchAnalysis.resume_id
        == resume_id
    ).delete(
        synchronize_session=False
    )

    # --------------------------------------------------------
    # Remove resume from applications
    # --------------------------------------------------------

    db.query(
        JobApplication
    ).filter(
        JobApplication.resume_version_id
        == resume_id
    ).update(
        {
            JobApplication.resume_version_id: None
        },
        synchronize_session=False,
    )

    # --------------------------------------------------------
    # Delete physical file
    # --------------------------------------------------------

    blob_url: str | None = None
    if resume.file_path and resume.file_path.startswith("vercel-blob:"):
        blob_url = _stored_blob_url(
            resume.file_path,
            user.id,
        )
    elif resume.file_path:

        try:

            file_path = (
                Path(
                    resume.file_path
                ).resolve()
            )

            upload_root = (
                UPLOAD_DIR.resolve()
            )

            try:

                file_path.relative_to(
                    upload_root
                )

                if file_path.is_file():

                    file_path.unlink()

            except ValueError:

                pass

        except (
            OSError,
            ValueError,
        ):

            pass

    # --------------------------------------------------------
    # Delete database record
    # --------------------------------------------------------

    try:

        db.delete(resume)

        db.commit()

    except Exception as exc:

        db.rollback()

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to delete resume: "
                f"{exc}"
            ),
        )

    if blob_url:
        try:
            delete_blob(blob_url, user.id)
        except HTTPException as exc:
            raise HTTPException(
                status_code=502,
                detail=(
                    "The resume record was deleted, but its stored file "
                    "could not be removed. Contact support to clean up "
                    "the orphaned file."
                ),
            ) from exc

    return {
        "message": (
            "Resume deleted successfully"
        ),
        "resume_id": resume_id,
    }


# ============================================================
# OPEN / DOWNLOAD RESUME
# ============================================================


@router.get(
    "/{resume_id}/file",
)
def open_resume_file(
    resume_id: int,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):

    resume = get_owned_resume(
        db,
        resume_id,
        user,
    )

    if not resume.file_path:

        raise HTTPException(
            status_code=404,
            detail=(
                "This resume does not have "
                "an uploaded file."
            ),
        )

    is_blob_file = resume.file_path.startswith("vercel-blob:")
    if is_blob_file:
        blob_url = _stored_blob_url(
            resume.file_path,
            user.id,
        )
        file_path = download_blob_to_temp(
            blob_url,
            user.id,
            resume.file_name or "resume",
        )
    else:
        file_path = validate_file_path(resume.file_path)

    if not file_path.is_file():
        if is_blob_file:
            file_path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=404,
            detail="Resume file is missing from storage.",
        )

    media_type = mimetypes.guess_type(
        resume.file_name
        or str(file_path)
    )[0]

    if not media_type:

        media_type = (
            "application/octet-stream"
        )

    return FileResponse(
        path=str(file_path),
        media_type=media_type,
        filename=(
            resume.file_name or file_path.name
        ),
        content_disposition_type="inline",
        background=(
            BackgroundTask(file_path.unlink, missing_ok=True)
            if is_blob_file
            else None
        ),
    )